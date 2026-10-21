#!/usr/bin/env python3
"""
testset_candidates.py - candidate items for the evaluation test set (v1).

WHAT THIS DOES
    Builds the working sheet from which the editor chooses the test items
    described in data/EVALUATION_TEST_SET_v1.md:

        test_candidates_<date>.xlsx   one row per candidate English sentence:
                                      the gold terms it carries (by string
                                      matching, every tag to be checked by
                                      hand), their gold renderings, any
                                      existing Kinyarwanda translation, the
                                      source and its licence, and empty
                                      columns for the editor's choice, the
                                      reference and the log of edits. A
                                      Coverage sheet counts candidates and kept
                                      items per gold term, and a suggested set
                                      marks a small selection that covers every
                                      term (CSV files instead when openpyxl is
                                      not installed).
        test_candidates_<date>.json   the same rows, with the run's inputs and
                                      their checksums.

    Candidates come from two kinds of input:
        --candidates FILE.json   sentences gathered from published sources,
                                 one object per sentence (term, sentence, url,
                                 page_title, publisher, licence, citation,
                                 note). English only.
        --bilingual FILE.docx    a document whose paragraphs give an English
                                 sentence followed by its Kinyarwanda
                                 translation. Every paragraph with a gold term
                                 becomes a candidate, with the existing
                                 translation beside it and the paragraph kept
                                 whole so the editor can check the split.

THE RULES IT APPLIES (data/EVALUATION_TEST_SET_v1.md)
    Gold: both reviewers' latest blind scores are 4, recomputed from the
    backup with scripts/agreement.py and checked against gold_terms.csv.
    The gold rendering is the Kinyarwanda string shown with those scores. A
    term whose rendering in the backup differs from it has left the gold: it
    is listed and never tagged.
    Tags: the English headword and its English variants, extended to plural
    and singular, verb forms of "To ..." headwords, an acronym given in
    brackets, hyphen or space between words, British spellings, and, for
    headwords of 12 letters or more, a hyphen or space inside a word
    (cephalo-pelvic). Every tag is checked by hand; the sheet only proposes.
    The dictionary's own example sentences are never test items: a candidate
    equal to, or containing, an example sentence from the backup (either
    language, ignoring case, spacing and punctuation) is dropped and listed.

GUARANTEES
    Reads only. Writes only into --out, never inside the repository: the
    candidates are research material held out from every system until the
    test file is frozen. Writes no Kinyarwanda of its own: every Kinyarwanda
    string in the output is copied from the backup or from a source document.

HOW TO RUN (repo root, venv active; openpyxl is optional)
    python scripts/testset_candidates.py \\
        --backup ~/LinguaMedica-backups/2026-10-17_1732 \\
        --candidates ~/Documents/LinguaMedica-testset/web_candidates.json \\
        --bilingual ~/Downloads/<history-taking script>.docx \\
        --out ~/Documents/LinguaMedica-testset
"""
import argparse
import csv
import hashlib
import importlib.util
import json
import re
import sys
import unicodedata
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

REPO_ROOT = Path(__file__).resolve().parent.parent
CORPUS_CSV = REPO_ROOT / "data" / "annotation_corpus_v1.csv"
CORPUS_DOC = REPO_ROOT / "data" / "ANNOTATION_CORPUS_v1.md"
GOLD_REVIEWERS = ("OU", "YV")
DOMAIN_ORDER = ("Diagnostics", "Infectious Disease", "Obstetrics", "Pharmacology")

_spec = importlib.util.spec_from_file_location("agreement", REPO_ROOT / "scripts" / "agreement.py")
agreement = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(agreement)


# --------------------------------------------------------------------------
# Inputs: the frozen corpus and the gold list
# --------------------------------------------------------------------------

def read_corpus():
    """{key: row} from the frozen corpus file, after checking its checksum."""
    recorded = re.search(r"SHA-256 \| `([0-9a-f]{64})`",
                         CORPUS_DOC.read_text(encoding="utf-8")).group(1)
    actual = hashlib.sha256(CORPUS_CSV.read_bytes()).hexdigest()
    if actual != recorded:
        sys.exit("data/annotation_corpus_v1.csv does not match the checksum recorded in "
                 "data/ANNOTATION_CORPUS_v1.md: the frozen corpus has been edited.")
    with open(CORPUS_CSV, encoding="utf-8", newline="") as f:
        return actual, {r["key"]: r for r in csv.DictReader(f)}


def read_gold(backup, corpus="v1", corpus_rows=None):
    """
    The gold terms of a closed-round backup.

    Returns (folder, manifest, checks, gold, left, examples): gold is
    {term_id: {...}} with the rendering shown to the reviewers; left lists the
    terms that were scored 4 by both but have left the gold; examples is the
    set of the dictionary's example sentences, normalised.
    """
    folder, manifest, checks, tables = agreement.read_backup(backup)
    terms = agreement.select_terms(tables["terms.csv"], corpus, "all")
    latest, _ = agreement.latest_blind(tables["term_reviews.csv"], set(terms), GOLD_REVIEWERS)
    a, b = GOLD_REVIEWERS
    gold, left = {}, []
    for term_id in sorted(set(latest[a]) & set(latest[b])):
        row_a, row_b = latest[a][term_id], latest[b][term_id]
        if int(row_a["score"]) != 4 or int(row_b["score"]) != 4:
            continue
        term = terms[term_id]
        shown = {(row_a.get("shown_rw") or "").strip(), (row_b.get("shown_rw") or "").strip()}
        reason = None
        if len(shown) != 1 or "" in shown:
            reason = "the two reviewers were not shown the same rendering"
        elif (term.get("kinyarwanda") or "").strip() != next(iter(shown)):
            reason = "the rendering changed after it was scored"
        if reason:
            left.append({"term_id": term_id, "english": term.get("english", ""), "reason": reason})
            continue
        key = (term.get("corpus_key") or "").strip()
        variants = ""
        if corpus_rows and key in corpus_rows:
            variants = corpus_rows[key].get("english_variants", "")
        gold[term_id] = {"term_id": term_id, "english": term.get("english", "").strip(),
                         "domain": (term.get("domain") or "").strip(), "corpus_key": key,
                         "rendering": next(iter(shown)), "english_variants": variants}
    listed = folder / "gold_terms.csv"
    if listed.is_file():
        with open(listed, encoding="utf-8", newline="") as f:
            ids = {int(r["term_id"]) for r in csv.DictReader(f)}
        if ids != set(gold) | {x["term_id"] for x in left}:
            sys.exit("gold_terms.csv in the backup does not match the gold recomputed from "
                     "its scores. Run scripts/agreement.py --terms all --round-closed again.")
    examples = set()
    for t in tables["terms.csv"]:
        for field in ("example_en", "example_rw"):
            if (t.get(field) or "").strip():
                examples.add(normalise(t[field]))
    return folder, manifest, checks, gold, left, examples


# --------------------------------------------------------------------------
# Tagging: string matching on the English side
# --------------------------------------------------------------------------

UK_SPELLINGS = (("haem", "hem"), ("aemia", "emia"), ("caesarean", "cesarean"),
                ("labour", "labor"), ("aetiology", "etiology"), ("oedema", "edema"),
                ("paediatric", "pediatric"), ("anaesthe", "anesthe"), ("oestro", "estro"))


def _british(phrase):
    forms = {phrase}
    for uk, us in UK_SPELLINGS:
        for form in list(forms):
            if us in form:
                forms.add(form.replace(us, uk))
    return forms


def _number_forms(word):
    """The word with its plural or singular (regular English rules, plus -fe/-ves)."""
    forms = {word}
    if word.endswith("ies"):
        forms.add(word[:-3] + "y")
    elif word.endswith("sis"):
        forms.add(word[:-3] + "ses")
    elif word.endswith("s") and not word.endswith(("ss", "us", "is")):
        forms.add(word[:-1])
    elif word.endswith("y") and len(word) > 1 and word[-2] not in "aeiou":
        forms.add(word[:-1] + "ies")
    elif re.search(r"(s|x|z|ch|sh)$", word):
        forms.add(word + "es")
    elif word.endswith("fe"):
        forms |= {word + "s", word[:-2] + "ves"}
    else:
        forms.add(word + "s")
    return forms


def _verb_forms(word):
    """Present, past and -ing forms, with and without a doubled last consonant."""
    stem = word[:-1] if word.endswith("e") else word
    forms = {word, word + "s", stem + "ed", stem + "ing"}
    if re.search(r"[^aeiou][aeiou][bdgklmnprt]$", word):
        forms |= {word + word[-1] + "ed", word + word[-1] + "ing"}
    return forms


def _phrase_regex(phrase):
    """Words joined by a space may also be joined by a hyphen, and the reverse."""
    parts = re.split(r"([ -])", phrase)
    out = []
    for part in parts:
        if part == " ":
            out.append(r"[\s-]")
        elif part == "-":
            out.append(r"[\s-]?")
        else:
            out.append(re.escape(part))
    return "".join(out)


def term_patterns(headword, variants=""):
    """(word patterns, acronyms) for one headword and its English variants."""
    names = [headword] + [v.strip() for v in (variants or "").split("/") if v.strip()]
    words, acronyms = set(), set()
    for name in names:
        acronyms |= set(re.findall(r"\(([A-Z]{2,})\)", name))
        base = re.sub(r"\s*\([^)]*\)", "", name).strip().lower()
        is_verb = base.startswith("to ")
        if is_verb:
            base = base[3:]
        base = re.sub(r"^(a|an|the) ", "", base)
        tokens = base.split()
        if not tokens:
            continue
        phrases = set()
        firsts = _verb_forms(tokens[0]) if is_verb else {tokens[0]}
        lasts = _number_forms(tokens[-1]) if (len(tokens) > 1 or not is_verb) else {None}
        for first in firsts:
            for last in lasts:
                if len(tokens) == 1:
                    phrases.add(first if last is None else last)
                else:
                    phrases.add(" ".join([first, *tokens[1:-1], last]))
        for phrase in phrases:
            for form in _british(phrase):
                words.add(_phrase_regex(form))
    return words, acronyms


def _compact(text):
    return re.sub(r"[\s-]", "", text.lower())


def _compact_forms(headword, variants=""):
    """Headword and variants with spaces and hyphens removed, if 12 letters or more."""
    out = set()
    for name in [headword] + [v.strip() for v in (variants or "").split("/") if v.strip()]:
        base = re.sub(r"\s*\([^)]*\)", "", name).strip().lower()
        base = re.sub(r"^(to|a|an|the) ", "", base)
        if len(_compact(base)) >= 12:
            out.add(_compact(base))
    return out


def compile_tags(gold):
    """[(gold entry, word regex, acronym regex or None, compact forms)], stable order."""
    compiled = []
    for entry in sorted(gold.values(), key=lambda g: g["english"].lower()):
        words, acronyms = term_patterns(entry["english"], entry.get("english_variants", ""))
        word_rx = re.compile(r"(?<![\w-])(" + "|".join(sorted(words)) + r")(?![\w-])", re.I)
        acr_rx = (re.compile(r"\b(" + "|".join(sorted(acronyms)) + r")s?\b")
                  if acronyms else None)
        compiled.append((entry, word_rx, acr_rx,
                         _compact_forms(entry["english"], entry.get("english_variants", ""))))
    return compiled


def tag(sentence, compiled):
    """[(gold entry, text as found)] for every gold term the sentence carries."""
    hits = []
    squeezed = _compact(sentence)
    for entry, word_rx, acr_rx, compacts in compiled:
        found = word_rx.search(sentence) or (acr_rx.search(sentence) if acr_rx else None)
        if found:
            hits.append((entry, found.group(0)))
        elif any(c in squeezed for c in compacts):
            hits.append((entry, "(written with a hyphen or space inside a word)"))
    return hits


# --------------------------------------------------------------------------
# Exclusion of the dictionary's example sentences
# --------------------------------------------------------------------------

def normalise(text):
    text = unicodedata.normalize("NFKC", text or "").lower()
    text = re.sub(r"[‘’`´]", "'", text)
    text = re.sub(r"[^\w\s']", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def is_dictionary_example(texts, examples):
    """True when any text equals a dictionary example or contains one of 6+ words."""
    for text in texts:
        norm = normalise(text)
        if not norm:
            continue
        if norm in examples:
            return True
        for example in examples:
            if len(example.split()) >= 6 and f" {example} " in f" {norm} ":
                return True
    return False


# --------------------------------------------------------------------------
# Bilingual documents: English sentence, then its Kinyarwanda translation
# --------------------------------------------------------------------------

ENGLISH_MARKERS = frozenset(
    "the and you your to of a an in on is was were are am did do does have has had any "
    "what how when where who which why this that these with for or not be been during "
    "it he she they we my his her their there if".split())


def docx_paragraphs(path):
    word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    paragraphs = []
    for p in root.iter(word + "p"):
        text = "".join(t.text or "" for t in p.iter(word + "t")).strip()
        if text:
            paragraphs.append(text)
    return paragraphs


def _english_share(text):
    tokens = re.findall(r"[A-Za-z']+", text.lower())
    if not tokens:
        return 0.0
    return sum(t in ENGLISH_MARKERS for t in tokens) / len(tokens)


def split_bilingual(paragraph):
    """
    (english, kinyarwanda) for a paragraph that gives an English sentence and
    then its translation. The split falls before the first clause after which
    no clause reads as English (English function words under 10%). A proposal
    only: the editor checks it against the paragraph.
    """
    clauses = [c for c in re.split(r"(?<=[?.!:])\s+", paragraph.strip()) if c]
    english_like = [_english_share(c) >= 0.10 for c in clauses]
    cut = len(clauses)
    for i in range(len(clauses)):
        if not any(english_like[i:]):
            cut = i
            break
    return " ".join(clauses[:cut]).strip(), " ".join(clauses[cut:]).strip()


def bilingual_candidates(path, compiled):
    paragraphs = docx_paragraphs(path)
    out = []
    i = 0
    while i < len(paragraphs):
        english, kinyarwanda = split_bilingual(paragraphs[i])
        original, location = paragraphs[i], f"paragraph {i + 1} (counting non-empty paragraphs)"
        if (english and not kinyarwanda and i + 1 < len(paragraphs)
                and _english_share(paragraphs[i + 1]) < 0.10):
            kinyarwanda = paragraphs[i + 1]
            original = paragraphs[i] + "\n" + paragraphs[i + 1]
            location = f"paragraphs {i + 1}-{i + 2} (counting non-empty paragraphs)"
            i += 1
        i += 1
        if not english or len(english.split()) < 4 or not tag(english, compiled):
            continue
        out.append({"sentence": english, "existing_rw": kinyarwanda, "original": original,
                    "source": Path(path).name, "location": location,
                    "licence": "unpublished document: ask its owner before publishing",
                    "kind": "bilingual document"})
    return out


# --------------------------------------------------------------------------
# Building the rows
# --------------------------------------------------------------------------

def item_id(sentence):
    return "C" + hashlib.sha256(normalise(sentence).encode("utf-8")).hexdigest()[:6].upper()


def build(gold, examples, web, bilingual):
    """(rows, coverage, left_out) from the gathered candidates."""
    compiled = compile_tags(gold)
    by_id, left_out = {}, []
    not_found = {}
    for c in web:
        if not (c.get("sentence") or "").strip():
            not_found.setdefault(c.get("term", ""), []).append(c.get("note", ""))
            continue
        c = dict(c, kind="published source")
        c.setdefault("existing_rw", "")
        c.setdefault("source", " - ".join(x for x in (c.get("publisher"), c.get("page_title")) if x))
        c.setdefault("location", c.get("url", ""))
        _add(c, by_id, left_out, compiled, examples)
    for c in bilingual:
        _add(c, by_id, left_out, compiled, examples)
    rows = list(by_id.values())
    chosen = set(suggest(rows))
    for r in rows:
        r["suggested"] = r["item"] in chosen
    rows.sort(key=_order)
    coverage = []
    for entry in sorted(gold.values(), key=lambda g: (_domain_rank(g["domain"]), g["english"].lower())):
        items = [r["item"] for r in rows if entry["english"] in r["terms"]]
        picked = [r["item"] for r in rows if r["suggested"] and entry["english"] in credited_terms(r)]
        coverage.append({"term": entry["english"], "domain": entry["domain"],
                         "rendering": entry["rendering"], "candidates": len(items),
                         "items": ", ".join(items), "suggested": ", ".join(picked),
                         "not_found_note": " | ".join(n for n in not_found.get(entry["english"], []) if n)})
    return rows, coverage, left_out


def credited_terms(row):
    """
    The terms a row is trusted to carry for the suggestion: for a bilingual
    document every tagged term; for a gathered sentence only the term it was
    gathered for and the terms its collector listed in it (the collector
    judged the sense; a tag the collector left out may be another sense).
    """
    if row["kind"] == "bilingual document":
        return set(row["terms"])
    return set(row["terms"]) & ({row.get("gathered_for", "")} | set(row.get("collector_terms", [])))


def suggest(rows):
    """
    Item ids of a small set of rows that covers every term the rows can cover:
    first one row per term from bilingual documents (their translation
    exists), then a greedy set cover, taking each time the row that adds the
    most uncovered terms. Ties: public domain, then CC BY, then 10 to 30
    words, then fewer words, then item id. A suggestion only.
    """
    def tiebreak(r):
        licence = r.get("licence", "")
        return ("public domain" not in licence.lower(), "CC BY" not in licence,
                not 10 <= r["words"] <= 30, r["words"], r["item"])
    chosen, covered = [], set()
    for r in sorted((r for r in rows if r["kind"] == "bilingual document"), key=tiebreak):
        if credited_terms(r) - covered:
            chosen.append(r["item"])
            covered |= credited_terms(r)
    uncovered = set().union(*(credited_terms(r) for r in rows)) - covered if rows else set()
    while uncovered:
        gain = max(len(credited_terms(r) & uncovered) for r in rows)
        if gain == 0:
            break
        pick = sorted((r for r in rows if len(credited_terms(r) & uncovered) == gain), key=tiebreak)[0]
        chosen.append(pick["item"])
        uncovered -= credited_terms(pick)
    return chosen


def _add(c, by_id, left_out, compiled, examples):
    sentence = c["sentence"].strip()
    if is_dictionary_example([sentence, c.get("existing_rw", "")], examples):
        left_out.append({"what": sentence, "why": "a dictionary example sentence"})
        return
    hits = tag(sentence, compiled)
    if not hits:
        left_out.append({"what": sentence, "why": f"no gold term matched (gathered for {c.get('term', '?')})"})
        return
    key = item_id(sentence)
    if key in by_id:
        return
    terms = [e["english"] for e, _ in hits]
    by_id[key] = {
        "item": key, "kind": c.get("kind", ""), "sentence": sentence,
        "terms": terms, "found_as": [found for _, found in hits],
        "renderings": [f"{e['english']}: {e['rendering']}" for e, _ in hits],
        "domain": hits[0][0]["domain"],
        "existing_rw": c.get("existing_rw", ""), "original": c.get("original", ""),
        "source": c.get("source", ""), "location": c.get("location", ""),
        "licence": c.get("licence", ""), "citation": c.get("citation", ""),
        "note": c.get("note", ""), "gathered_for": c.get("term", ""),
        "collector_terms": list(c.get("other_terms") or []),
        "words": len(sentence.split()), "suggested": False,
    }


def _domain_rank(domain):
    return DOMAIN_ORDER.index(domain) if domain in DOMAIN_ORDER else len(DOMAIN_ORDER)


def _order(row):
    return (_domain_rank(row["domain"]), not row.get("suggested"), row["terms"][0].lower(),
            -len(row["terms"]), row["item"])


# --------------------------------------------------------------------------
# Output
# --------------------------------------------------------------------------

COLUMNS = (
    ("item", "Item", 9), ("keep", "Keep?", 9), ("suggested", "Suggested", 10),
    ("terms", "Gold terms (check each)", 24),
    ("sentence", "English sentence (verbatim)", 60), ("found_as", "Found as", 18),
    ("renderings", "Gold renderings", 34),
    ("wrong_sense", "Tagged terms used in another sense (type them)", 22),
    ("existing_rw", "Existing Kinyarwanda translation", 46),
    ("original", "Original paragraph (bilingual documents: check the split)", 46),
    ("reference", "Reference (Kinyarwanda)", 46),
    ("reference_by", "Reference by", 16), ("edits", "Edits to the existing translation", 34),
    ("source", "Source", 30), ("location", "URL or place in the document", 34),
    ("licence", "Licence", 20), ("citation", "Citation", 30), ("note", "Note", 30),
    ("words", "Words", 7), ("tags", "_tags", 4),
)

READ_ME = (
    "What this sheet is: candidate English sentences for the evaluation test set (data/EVALUATION_TEST_SET_v1.md), built by scripts/testset_candidates.py. Nothing here is a test item until you keep it.",
    "1. Suggested: a small set that covers every gold term the candidates can cover, picked by the script (one existing translation per term first, then the sentences that add the most uncovered terms). Start there; it is only a suggestion.",
    "2. Keep? Keep, Maybe or Drop. Aim for every gold term in at least one kept item; a sentence with several gold terms counts for each.",
    "3. Gold terms: proposed by string matching. Check every one. If a tagged term is used in another sense (a statistical test, delivery of oxygen), type it in 'Tagged terms used in another sense'; Coverage then stops counting it for that row.",
    "4. Reference: where an existing human translation is given, start from it, give each tagged term its gold rendering, and write every change in 'Edits to the existing translation'. Where there is none, translate the sentence yourself and set 'Reference by' to Editor (CM).",
    "5. Open the URL before you translate: the sentence was quoted from it, and reading it in context also helps the translation.",
    "6. Licence: only sentences you may publish go into the frozen test file. 'Unpublished document' rows need their owner's permission first.",
    "7. Coverage shows, per gold term, how many candidates there are, which are suggested, and how many you have kept (wrong-sense rows excluded).",
    "This sheet stays outside the repository and away from every system until the test file is frozen (SHA-256 and git tag).",
)


def _cell(row, key):
    value = row.get(key, "")
    if key in ("terms", "found_as"):
        return "; ".join(value)
    if key == "renderings":
        return "\n".join(value)
    if key == "tags":
        return "|" + "|".join(row["terms"]) + "|"
    if key == "suggested":
        return "Yes" if row.get("suggested") else ""
    return value


def write_outputs(rows, coverage, left_out, run, out_dir, date, force=False):
    out_dir = agreement._outside_repo(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = f"test_candidates_{date}"
    existing = [p for p in out_dir.glob(f"{stem}*") if p.is_file()]
    if existing and not force:
        sys.exit(f"{existing[0].name} already exists in {out_dir}: the editor's choices and "
                 "references live in that sheet, so it is never overwritten. Use another --date "
                 "(or --force for a sheet nobody has worked in).")
    (out_dir / f"{stem}.json").write_text(
        json.dumps({"run": run, "rows": rows, "coverage": coverage, "left_out": left_out},
                   ensure_ascii=False, indent=1), encoding="utf-8")
    try:
        import openpyxl
    except ImportError:
        _write_csvs(rows, coverage, out_dir, stem)
        return [out_dir / f"{stem}.json", out_dir / f"{stem}.csv", out_dir / f"{stem}_coverage.csv"]
    _write_xlsx(openpyxl, rows, coverage, left_out, out_dir / f"{stem}.xlsx")
    return [out_dir / f"{stem}.json", out_dir / f"{stem}.xlsx"]


def _write_csvs(rows, coverage, out_dir, stem):
    with open(out_dir / f"{stem}.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow([label for _, label, _ in COLUMNS])
        for r in rows:
            w.writerow([_cell(r, key) for key, _, _ in COLUMNS])
    with open(out_dir / f"{stem}_coverage.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Gold term", "Domain", "Gold rendering", "Candidates", "Suggested", "Items", "Not found"])
        for c in coverage:
            w.writerow([c["term"], c["domain"], c["rendering"], c["candidates"], c["suggested"],
                        c["items"], c["not_found_note"]])


def _write_xlsx(openpyxl, rows, coverage, left_out, path):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="F3E7E7")
    wrap = Alignment(wrap_text=True, vertical="top")

    ws = wb.active
    ws.title = "Candidates"
    ws.append([label for _, label, _ in COLUMNS])
    for r in rows:
        ws.append([_cell(r, key) for key, _, _ in COLUMNS])
    n = len(rows) + 1
    for i, (key, _, width) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = width
        if key == "tags":
            ws.column_dimensions[letter].hidden = True
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment = Font(bold=True), head_fill, wrap
    for line in ws.iter_rows(min_row=2, max_row=n):
        for cell in line:
            cell.alignment = wrap
    ws.freeze_panes = "D2"
    col = {key: get_column_letter(i) for i, (key, _, _) in enumerate(COLUMNS, start=1)}
    for key, options in (("keep", "Keep,Maybe,Drop"),
                         ("reference_by", "Existing translation,Editor (CM)")):
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col[key]}2:{col[key]}{max(n, 2)}")

    cv = wb.create_sheet("Coverage")
    cv.append(["Gold term", "Domain", "Gold rendering", "Candidates", "Kept", "Suggested",
               "All items", "Not found"])
    for i, c in enumerate(coverage, start=2):
        keep = f'Candidates!${col["keep"]}:${col["keep"]},"Keep"'
        tags = f'Candidates!${col["tags"]}:${col["tags"]},"*|"&A{i}&"|*"'
        wrong = f'Candidates!${col["wrong_sense"]}:${col["wrong_sense"]},"*"&A{i}&"*"'
        kept = f"=COUNTIFS({keep},{tags})-COUNTIFS({keep},{tags},{wrong})"
        cv.append([c["term"], c["domain"], c["rendering"], c["candidates"], kept, c["suggested"],
                   c["items"], c["not_found_note"]])
    for letter, width in zip("ABCDEFGH", (28, 18, 36, 11, 8, 22, 40, 50)):
        cv.column_dimensions[letter].width = width
    for cell in cv[1]:
        cell.font, cell.fill = Font(bold=True), head_fill
    cv.freeze_panes = "B2"

    lo = wb.create_sheet("Left out")
    lo.append(["Left out", "Why"])
    for x in left_out:
        lo.append([x["what"], x["why"]])
    lo.column_dimensions["A"].width, lo.column_dimensions["B"].width = 90, 50
    for cell in lo[1]:
        cell.font, cell.fill = Font(bold=True), head_fill

    rm = wb.create_sheet("Read me")
    for line in READ_ME:
        rm.append([line])
    rm.column_dimensions["A"].width = 140
    for line in rm.iter_rows():
        for cell in line:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def _sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main(argv=None):
    parser = argparse.ArgumentParser(description="Candidate items for the evaluation test set.")
    parser.add_argument("--backup", required=True, help="a closed-round backup folder")
    parser.add_argument("--candidates", action="append", default=[],
                        help="JSON list of sentences gathered from sources (repeatable)")
    parser.add_argument("--bilingual", action="append", default=[],
                        help="a .docx with English then Kinyarwanda in each paragraph (repeatable)")
    parser.add_argument("--out", required=True, help="output folder, never inside the repository")
    parser.add_argument("--corpus", default="v1", help="corpus version (default v1)")
    parser.add_argument("--date", default=None, help="date in the file names (default today)")
    parser.add_argument("--force", action="store_true",
                        help="overwrite a sheet of the same date (only one nobody has worked in)")
    args = parser.parse_args(argv)

    corpus_sha, corpus_rows = read_corpus()
    folder, manifest, checks, gold, left, examples = read_gold(args.backup, args.corpus, corpus_rows)
    web = []
    for path in args.candidates:
        web += json.loads(Path(path).expanduser().read_text(encoding="utf-8"))
    compiled = compile_tags(gold)
    bilingual = []
    for path in args.bilingual:
        bilingual += bilingual_candidates(Path(path).expanduser(), compiled)
    rows, coverage, left_out = build(gold, examples, web, bilingual)
    left_out = [{"what": x["english"], "why": x["reason"]} for x in left] + left_out
    run = {
        "script": "scripts/testset_candidates.py", "script_sha256": _sha256(Path(__file__)),
        "ran_at": datetime.now().astimezone().isoformat(),
        "repo": agreement.repo_state(), "python": sys.version.split()[0],
        "backup": folder.name, "backup_checksums": checks, "corpus_sha256": corpus_sha,
        "inputs": {str(Path(p).expanduser().name): _sha256(Path(p).expanduser())
                   for p in [*args.candidates, *args.bilingual]},
        "gold_terms": len(gold), "left_the_gold": len(left), "candidates": len(rows),
    }
    date = args.date or datetime.now().astimezone().date().isoformat()
    written = write_outputs(rows, coverage, left_out, run, args.out, date, force=args.force)
    covered = sum(1 for c in coverage if c["candidates"])
    print(f"{len(rows)} candidate sentences; {covered} of {len(gold)} gold terms have at least one.")
    missing = [c["term"] for c in coverage if not c["candidates"]]
    if missing:
        print("No candidate yet: " + ", ".join(missing))
    if left:
        print("Left the gold: " + ", ".join(x["english"] for x in left))
    for path in written:
        print(f"Wrote {path}")


if __name__ == "__main__":
    main()
